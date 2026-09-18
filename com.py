#!/usr/bin/env python3
"""Estimate and validate the Consumer Card fee revenue model.

All coefficients, diagnostics, accuracy statistics, and stability results are
calculated from the source workbooks at run time. No benchmark model results
from an existing report are embedded in this file.

The default run writes full-precision CSV tables, a JSON manifest, figures,
and an English Word report. Use --skip-report to omit the DOCX.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import sys
import warnings
from dataclasses import dataclass
from importlib import metadata
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from scipy import stats
from statsmodels.stats.diagnostic import breaks_cusumolsresid
from statsmodels.stats.outliers_influence import variance_inflation_factor
from statsmodels.stats.stattools import durbin_watson, jarque_bera
from statsmodels.tsa.stattools import adfuller, kpss


ROOT = Path(__file__).resolve().parent
DEFAULT_CARD = Path("/Users/hyx/Desktop/Cardfee/Cardfee.xlsx")
DEFAULT_MACRO = Path("/Users/hyx/Desktop/Cardfee/Revise Version/Macro_Data_201001-202612.xls")
DEFAULT_OUTPUT = ROOT / "output" / "consumer_model_clean"

IN_SAMPLE_START = pd.Timestamp("2011-01-01")
IN_SAMPLE_END = pd.Timestamp("2023-09-01")
OOT_START = pd.Timestamp("2023-10-01")
OOT_END = pd.Timestamp("2025-12-01")

DATE_COLUMN = "YYYYMM"
TARGET_COLUMN = "Consumer Card"

MACRO_SPECS = {
    "Unemployment": {
        "group": "Core",
        "source": "M_FLBR_B.IUSA",
        "transform": "Monthly first difference in percentage points",
    },
    "CPI": {
        "group": "Core",
        "source": "M_FCPIU_B.IUSA",
        "transform": "Monthly percentage change",
    },
    "DPI": {
        "group": "Core",
        "source": "M_FYPDPIQ_B.IUSA",
        "transform": "Monthly percentage change",
    },
    "RetailSales": {
        "group": "Consumer",
        "source": "M_FRT_B.IUSA (Retail Sales Total_Bil.USD,CDASAAR)",
        "transform": "Monthly percentage change",
    },
    "ConsumerConfidence": {
        "group": "Consumer",
        "source": "M_FCBC_B.IUSA(Consumer confidence index)",
        "transform": "Monthly percentage change",
    },
}

BASE_TERMS = ["y_lag1", "y_lag2", "January", "December"]
SELECTION_P_LIMIT = 0.05
VIF_LIMIT = 10.0
DW_LOWER = 1.5
DW_UPPER = 2.5
JB_P_LIMIT = 0.05
ROLLING_WINDOW = 60


@dataclass
class Results:
    data: pd.DataFrame
    in_sample: pd.DataFrame
    oot: pd.DataFrame
    model: object
    selected_terms: list[str]
    stationarity: pd.DataFrame
    screening: pd.DataFrame
    coefficients: pd.DataFrame
    predictions: pd.DataFrame
    accuracy: pd.DataFrame
    residual_diagnostics: pd.DataFrame
    cusum: pd.DataFrame
    chow: pd.DataFrame
    rolling: pd.DataFrame
    rolling_summary: pd.DataFrame
    oot_only: pd.DataFrame
    source_audit: dict


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_month(values: pd.Series) -> pd.DatetimeIndex:
    if pd.api.types.is_datetime64_any_dtype(values):
        parsed = pd.to_datetime(values)
    else:
        text = values.astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
        parsed = pd.to_datetime(text, format="%Y%m", errors="coerce")
        if parsed.isna().any():
            parsed = pd.to_datetime(values, errors="coerce")
    if parsed.isna().any():
        raise ValueError("The date column contains values that cannot be parsed as monthly dates.")
    return pd.DatetimeIndex(parsed).to_period("M").to_timestamp()


def validate_calendar(frame: pd.DataFrame, name: str) -> dict:
    if frame.index.has_duplicates:
        duplicates = frame.index[frame.index.duplicated()].strftime("%Y-%m").tolist()
        raise ValueError(f"{name} contains duplicate months: {duplicates}")
    expected = pd.date_range(frame.index.min(), frame.index.max(), freq="MS")
    missing = expected.difference(frame.index)
    if len(missing):
        raise ValueError(f"{name} is missing months: {missing.strftime('%Y-%m').tolist()}")
    return {
        "rows": int(len(frame)),
        "start": frame.index.min().strftime("%Y-%m"),
        "end": frame.index.max().strftime("%Y-%m"),
        "duplicate_months": 0,
        "missing_calendar_months": 0,
    }


def read_sources(card_path: Path, macro_path: Path) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    if not card_path.is_file():
        raise FileNotFoundError(f"Card source not found: {card_path}")
    if not macro_path.is_file():
        raise FileNotFoundError(f"Macro source not found: {macro_path}")

    card = pd.read_excel(card_path, sheet_name="汇总数据")
    missing_card = {DATE_COLUMN, TARGET_COLUMN}.difference(card.columns)
    if missing_card:
        raise ValueError(f"Card source is missing columns: {sorted(missing_card)}")
    card = card[[DATE_COLUMN, TARGET_COLUMN]].copy()
    card.index = normalize_month(card.pop(DATE_COLUMN))
    card.index.name = "date"
    card[TARGET_COLUMN] = pd.to_numeric(card[TARGET_COLUMN], errors="coerce")

    macro = pd.read_excel(macro_path, sheet_name="Macro Data", header=1)
    macro_fields = [spec["source"] for spec in MACRO_SPECS.values()]
    missing_macro = {DATE_COLUMN, *macro_fields}.difference(macro.columns)
    if missing_macro:
        raise ValueError(f"Macro source is missing columns: {sorted(missing_macro)}")
    macro = macro[[DATE_COLUMN, *macro_fields]].copy()
    macro.index = normalize_month(macro.pop(DATE_COLUMN))
    macro.index.name = "date"
    for column in macro.columns:
        macro[column] = pd.to_numeric(macro[column], errors="coerce")

    card_audit = validate_calendar(card, "Card data")
    macro_audit = validate_calendar(macro, "Macro data")
    if card[TARGET_COLUMN].isna().any():
        raise ValueError("Consumer Card contains missing values.")
    if (card[TARGET_COLUMN] <= 0).any():
        raise ValueError("Consumer Card must be positive for the log transformation.")
    if macro[macro_fields].isna().any().any():
        counts = macro[macro_fields].isna().sum()
        raise ValueError(f"Macro data contain missing values: {counts[counts > 0].to_dict()}")

    return card, macro, {
        "card_filename": card_path.name,
        "macro_filename": macro_path.name,
        "card_path": str(card_path),
        "macro_path": str(macro_path),
        "card_sha256": sha256(card_path),
        "macro_sha256": sha256(macro_path),
        "card": card_audit,
        "macro": macro_audit,
    }


def engineer_data(card: pd.DataFrame, macro: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    data = card.join(macro, how="left")
    modeling_range = data.loc[IN_SAMPLE_START:OOT_END]
    if modeling_range.empty:
        raise ValueError("The modeling period is outside the source data range.")
    macro_fields = [spec["source"] for spec in MACRO_SPECS.values()]
    if modeling_range[macro_fields].isna().any().any():
        raise ValueError("Macro values are missing inside the modeling period.")

    data["y"] = 100.0 * np.log(data[TARGET_COLUMN]).diff()
    data["y_lag1"] = data["y"].shift(1)
    data["y_lag2"] = data["y"].shift(2)
    data["January"] = (data.index.month == 1).astype(int)
    data["December"] = (data.index.month == 12).astype(int)

    candidates: list[str] = []
    for name, spec in MACRO_SPECS.items():
        source = spec["source"]
        transformed = data[source].diff() if name == "Unemployment" else data[source].pct_change(fill_method=None) * 100.0
        current_name = f"{name}_chg"
        average_name = f"{current_name}_MA3"
        data[current_name] = transformed
        data[average_name] = transformed.rolling(3).mean()
        candidates.extend([current_name, average_name])

    expected = pd.date_range(IN_SAMPLE_START, OOT_END, freq="MS")
    missing = expected.difference(data.index)
    if len(missing):
        raise ValueError(f"Combined data are missing months: {missing.strftime('%Y-%m').tolist()}")
    return data, candidates


def fit_ols(frame: pd.DataFrame, terms: list[str]):
    sample = frame[["y", *terms]].dropna().copy()
    if len(sample) <= len(terms) + 1:
        raise ValueError("Insufficient observations for the requested regression.")
    x = sm.add_constant(sample[terms], has_constant="add")
    return sm.OLS(sample["y"], x).fit(), sample


def compute_vifs(model) -> dict[str, float]:
    x = pd.DataFrame(model.model.exog, columns=model.model.exog_names)
    return {
        name: float(variance_inflation_factor(x.values, index))
        for index, name in enumerate(x.columns)
        if name != "const"
    }


def model_diagnostics(model) -> dict[str, float]:
    jb = jarque_bera(model.resid)
    return {
        "n": int(model.nobs),
        "r2": float(model.rsquared),
        "adj_r2": float(model.rsquared_adj),
        "dw": float(durbin_watson(model.resid)),
        "jb_stat": float(jb[0]),
        "jb_p": float(jb[1]),
    }


def forward_selection(data: pd.DataFrame, candidates: list[str]) -> tuple[list[str], pd.DataFrame]:
    sample = data.loc[IN_SAMPLE_START:IN_SAMPLE_END]
    selected: list[str] = []
    remaining = list(candidates)
    records: list[dict] = []
    round_number = 1

    while remaining:
        trials = []
        for candidate in remaining:
            terms = BASE_TERMS + selected + [candidate]
            model, used = fit_ols(sample, terms)
            vifs = compute_vifs(model)
            diagnostics = model_diagnostics(model)
            trial = {
                "round": round_number,
                "candidate": candidate,
                "n": int(len(used)),
                "coefficient": float(model.params[candidate]),
                "p_value": float(model.pvalues[candidate]),
                "adj_r2": float(model.rsquared_adj),
                "dw": diagnostics["dw"],
                "jb_p": diagnostics["jb_p"],
                "max_vif": float(max(vifs.values())),
            }
            trial["passes"] = bool(
                trial["p_value"] < SELECTION_P_LIMIT
                and trial["max_vif"] < VIF_LIMIT
                and DW_LOWER <= trial["dw"] <= DW_UPPER
                and trial["jb_p"] > JB_P_LIMIT
            )
            trials.append(trial)
            records.append(trial.copy())

        passing = [trial for trial in trials if trial["passes"]]
        if not passing:
            break
        winner = min(passing, key=lambda item: (item["p_value"], -item["adj_r2"], item["candidate"]))
        selected.append(winner["candidate"])
        remaining.remove(winner["candidate"])
        round_number += 1

    return selected, pd.DataFrame(records)


def stationarity_table(data: pd.DataFrame) -> pd.DataFrame:
    level = data.loc[IN_SAMPLE_START:IN_SAMPLE_END, TARGET_COLUMN]
    difference = level.diff().dropna()
    rows = []
    for label, series in (("Consumer Card level", level), ("Consumer Card first difference", difference)):
        adf = adfuller(series, regression="c", autolag="AIC", result_object=False)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            kp = kpss(series, regression="c", nlags="auto")
        rows.append(
            {
                "series": label,
                "n": int(len(series)),
                "adf_stat": float(adf[0]),
                "adf_p": float(adf[1]),
                "adf_lags": int(adf[2]),
                "kpss_stat": float(kp[0]),
                "kpss_p": float(kp[1]),
                "kpss_lags": int(kp[2]),
            }
        )
    return pd.DataFrame(rows)


def coefficient_table(model) -> pd.DataFrame:
    vifs = compute_vifs(model)
    confidence = model.conf_int()
    return pd.DataFrame(
        [
            {
                "term": term,
                "coefficient": float(model.params[term]),
                "std_error": float(model.bse[term]),
                "t_stat": float(model.tvalues[term]),
                "p_value": float(model.pvalues[term]),
                "ci_low": float(confidence.loc[term, 0]),
                "ci_high": float(confidence.loc[term, 1]),
                "vif": vifs.get(term, np.nan),
            }
            for term in model.params.index
        ]
    )


def prediction_frame(data: pd.DataFrame, model, sample: pd.DataFrame, window: str) -> pd.DataFrame:
    terms = [name for name in model.params.index if name != "const"]
    x = sm.add_constant(sample[terms], has_constant="add")
    predicted_y = model.predict(x)
    prior_level = data[TARGET_COLUMN].shift(1).reindex(sample.index)
    predicted_level = prior_level * np.exp(predicted_y / 100.0)
    actual_level = data[TARGET_COLUMN].reindex(sample.index)
    return pd.DataFrame(
        {
            "window": window,
            "actual_level": actual_level,
            "predicted_level": predicted_level,
            "actual_y": sample["y"],
            "predicted_y": predicted_y,
            "residual_y": sample["y"] - predicted_y,
        },
        index=sample.index,
    )


def accuracy_row(predictions: pd.DataFrame) -> dict:
    error = predictions["actual_level"] - predictions["predicted_level"]
    return {
        "window": predictions["window"].iloc[0],
        "n": int(len(predictions)),
        "mape_pct": float((error.abs() / predictions["actual_level"]).mean() * 100.0),
        "rmse": float(np.sqrt(np.mean(np.square(error)))),
        "mae": float(np.mean(np.abs(error))),
    }


def residual_row(predictions: pd.DataFrame) -> dict:
    residual = predictions["residual_y"]
    jb = jarque_bera(residual)
    return {
        "window": predictions["window"].iloc[0],
        "n": int(len(predictions)),
        "mean": float(residual.mean()),
        "std": float(residual.std(ddof=1)),
        "dw": float(durbin_watson(residual)),
        "jb_p": float(jb[1]),
    }


def cusum_table(data: pd.DataFrame, terms: list[str]) -> pd.DataFrame:
    rows = []
    for label, end in (("In sample", IN_SAMPLE_END), ("Full sample", OOT_END)):
        model, sample = fit_ols(data.loc[IN_SAMPLE_START:end], terms)
        statistic, p_value, critical = breaks_cusumolsresid(model.resid, ddof=len(model.params))
        rows.append(
            {
                "sample": label,
                "n": int(len(sample)),
                "statistic": float(statistic),
                "p_value": float(p_value),
                "critical_1pct": float(critical[0][1]),
                "critical_5pct": float(critical[1][1]),
                "critical_10pct": float(critical[2][1]),
            }
        )
    return pd.DataFrame(rows)


def chow_table(data: pd.DataFrame, terms: list[str]) -> pd.DataFrame:
    sample = data.loc[IN_SAMPLE_START:OOT_END, ["y", *terms]].dropna()
    pooled, _ = fit_ols(sample, terms)
    parameter_count = len(pooled.params)
    break_dates = (OOT_START, OOT_START + pd.DateOffset(months=3), OOT_START + pd.DateOffset(months=9))
    rows = []
    for break_date in break_dates:
        pre = sample.loc[sample.index < break_date]
        post = sample.loc[sample.index >= break_date]
        pre_model, _ = fit_ols(pre, terms)
        post_model, _ = fit_ols(post, terms)
        numerator = (pooled.ssr - pre_model.ssr - post_model.ssr) / parameter_count
        denominator_df = len(pre) + len(post) - 2 * parameter_count
        denominator = (pre_model.ssr + post_model.ssr) / denominator_df
        f_stat = float(numerator / denominator)
        rows.append(
            {
                "break_date": break_date.strftime("%Y-%m-%d"),
                "n_pre": int(len(pre)),
                "n_post": int(len(post)),
                "f_stat": f_stat,
                "p_value": float(stats.f.sf(f_stat, parameter_count, denominator_df)),
                "df_num": int(parameter_count),
                "df_den": int(denominator_df),
            }
        )
    return pd.DataFrame(rows)


def rolling_coefficients(data: pd.DataFrame, terms: list[str]) -> tuple[pd.DataFrame, pd.DataFrame]:
    sample = data.loc[IN_SAMPLE_START:OOT_END, ["y", *terms]].dropna()
    if len(sample) < ROLLING_WINDOW:
        raise ValueError("The usable model sample is shorter than the rolling window.")
    rows = []
    for end_position in range(ROLLING_WINDOW - 1, len(sample)):
        current = sample.iloc[end_position - ROLLING_WINDOW + 1 : end_position + 1]
        model, _ = fit_ols(current, terms)
        row = {
            "window_end": current.index[-1],
            "window_start": current.index[0],
            "n": int(len(current)),
        }
        row.update({term: float(model.params[term]) for term in model.params.index})
        rows.append(row)
    rolling = pd.DataFrame(rows).set_index("window_end")

    summary = []
    for term in terms:
        values = rolling[term]
        initial_sign = np.sign(values.iloc[0])
        summary.append(
            {
                "term": term,
                "min": float(values.min()),
                "max": float(values.max()),
                "mean": float(values.mean()),
                "std": float(values.std(ddof=1)),
                "same_sign_count": int((np.sign(values) == initial_sign).sum()),
                "windows": int(len(values)),
            }
        )
    return rolling, pd.DataFrame(summary)


def run_analysis(card_path: Path, macro_path: Path) -> Results:
    card, macro, source_audit = read_sources(card_path, macro_path)
    data, candidates = engineer_data(card, macro)
    selected_macro, screening = forward_selection(data, candidates)
    if not selected_macro:
        raise RuntimeError("Forward selection did not retain a macro candidate under the configured gates.")
    selected_terms = BASE_TERMS + selected_macro

    model, in_sample = fit_ols(data.loc[IN_SAMPLE_START:IN_SAMPLE_END], selected_terms)
    _, oot = fit_ols(data.loc[OOT_START:OOT_END], selected_terms)
    predictions_is = prediction_frame(data, model, in_sample, "In sample")
    predictions_oot = prediction_frame(data, model, oot, "Out of time")
    predictions = pd.concat([predictions_is, predictions_oot]).sort_index()
    rolling, rolling_summary = rolling_coefficients(data, selected_terms)
    oot_model, oot_sample = fit_ols(data.loc[OOT_START:OOT_END], selected_terms)
    oot_only = coefficient_table(oot_model)
    oot_only["n"] = len(oot_sample)

    return Results(
        data=data,
        in_sample=in_sample,
        oot=oot,
        model=model,
        selected_terms=selected_terms,
        stationarity=stationarity_table(data),
        screening=screening,
        coefficients=coefficient_table(model),
        predictions=predictions,
        accuracy=pd.DataFrame([accuracy_row(predictions_is), accuracy_row(predictions_oot)]),
        residual_diagnostics=pd.DataFrame([residual_row(predictions_is), residual_row(predictions_oot)]),
        cusum=cusum_table(data, selected_terms),
        chow=chow_table(data, selected_terms),
        rolling=rolling,
        rolling_summary=rolling_summary,
        oot_only=oot_only,
        source_audit=source_audit,
    )


def save_outputs(results: Results, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    frames = {
        "stationarity.csv": results.stationarity,
        "macro_screening.csv": results.screening,
        "coefficients.csv": results.coefficients,
        "predictions.csv": results.predictions.reset_index(),
        "accuracy.csv": results.accuracy,
        "residual_diagnostics.csv": results.residual_diagnostics,
        "cusum.csv": results.cusum,
        "chow.csv": results.chow,
        "rolling_coefficients.csv": results.rolling.reset_index(),
        "rolling_summary.csv": results.rolling_summary,
        "oot_only_coefficients.csv": results.oot_only,
    }
    for filename, frame in frames.items():
        frame.to_csv(output_dir / filename, index=False, encoding="utf-8-sig", float_format="%.12g")

    manifest = {
        "periods": {
            "in_sample": [IN_SAMPLE_START.strftime("%Y-%m"), IN_SAMPLE_END.strftime("%Y-%m")],
            "out_of_time": [OOT_START.strftime("%Y-%m"), OOT_END.strftime("%Y-%m")],
        },
        "target": TARGET_COLUMN,
        "selected_terms": results.selected_terms,
        "selection_gates": {
            "p_value_below": SELECTION_P_LIMIT,
            "vif_below": VIF_LIMIT,
            "durbin_watson_range": [DW_LOWER, DW_UPPER],
            "jarque_bera_p_above": JB_P_LIMIT,
        },
        "model_diagnostics": model_diagnostics(results.model),
        "accuracy": results.accuracy.to_dict(orient="records"),
        "source_audit": results.source_audit,
        "runtime": {
            "python": sys.version,
            "platform": platform.platform(),
            "packages": {
                name: metadata.version(name)
                for name in ("numpy", "pandas", "scipy", "statsmodels", "matplotlib", "python-docx")
            },
        },
    }
    (output_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def make_plots(results: Results, assets_dir: Path) -> dict[str, Path]:
    assets_dir.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 9.5,
            "axes.titlesize": 11,
            "axes.labelsize": 9.5,
            "axes.edgecolor": "#AAB2BD",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "grid.color": "#E5E7EB",
            "grid.linewidth": 0.7,
        }
    )

    history_path = assets_dir / "consumer_card_history.png"
    fig, ax = plt.subplots(figsize=(7.2, 3.2), dpi=180)
    history = results.data.loc[IN_SAMPLE_START:OOT_END, TARGET_COLUMN]
    ax.plot(history.index, history, color="#24476B", linewidth=1.6)
    ax.axvspan(OOT_START, OOT_END, color="#DCEAF4", alpha=0.8, label="Out of time")
    ax.axvline(OOT_START, color="#6B7280", linewidth=1.0, linestyle="--")
    ax.set_title("Consumer Card Fee Revenue")
    ax.set_ylabel("Source units")
    ax.grid(axis="y")
    ax.legend(frameon=False, loc="upper right")
    ax.xaxis.set_major_locator(mdates.YearLocator(2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    fig.tight_layout()
    fig.savefig(history_path, bbox_inches="tight")
    plt.close(fig)

    fit_path = assets_dir / "actual_vs_predicted.png"
    fig, ax = plt.subplots(figsize=(7.2, 3.35), dpi=180)
    actual = results.data.loc[results.in_sample.index.min():OOT_END, TARGET_COLUMN]
    ax.plot(actual.index, actual, color="#111827", linewidth=1.45, label="Actual")
    in_pred = results.predictions.loc[results.predictions["window"] == "In sample", "predicted_level"]
    oot_pred = results.predictions.loc[results.predictions["window"] == "Out of time", "predicted_level"]
    ax.plot(in_pred.index, in_pred, color="#4F7CAC", linewidth=1.25, label="In sample fitted")
    ax.plot(oot_pred.index, oot_pred, color="#C76D3A", linewidth=1.5, label="OOT one step")
    ax.axvline(OOT_START, color="#6B7280", linewidth=1.0, linestyle="--")
    ax.set_title("Actual and One Step Predicted Levels")
    ax.set_ylabel("Source units")
    ax.grid(axis="y")
    ax.legend(frameon=False, ncol=3, loc="upper right")
    ax.xaxis.set_major_locator(mdates.YearLocator(2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    fig.tight_layout()
    fig.savefig(fit_path, bbox_inches="tight")
    plt.close(fig)

    rolling_path = assets_dir / "rolling_coefficients.png"
    macro_terms = [term for term in results.selected_terms if term not in BASE_TERMS]
    fig, axes = plt.subplots(2, 1, figsize=(7.2, 5.0), dpi=180, sharex=True)
    axes[0].plot(results.rolling.index, results.rolling["y_lag1"], label="y lag 1", color="#24476B")
    axes[0].plot(results.rolling.index, results.rolling["y_lag2"], label="y lag 2", color="#4F7CAC")
    colors = ["#C76D3A", "#7C3AED", "#0F766E", "#B45309"]
    for index, term in enumerate(macro_terms):
        axes[0].plot(results.rolling.index, results.rolling[term], label=term, color=colors[index % len(colors)])
    axes[0].axhline(0, color="#9CA3AF", linewidth=0.8)
    axes[0].set_title(f"Rolling {ROLLING_WINDOW} Month Dynamic and Macro Coefficients")
    axes[0].legend(frameon=False, ncol=3, loc="upper right")
    axes[0].grid(axis="y")
    axes[1].plot(results.rolling.index, results.rolling["January"], label="January", color="#7C3AED")
    axes[1].plot(results.rolling.index, results.rolling["December"], label="December", color="#0F766E")
    axes[1].axhline(0, color="#9CA3AF", linewidth=0.8)
    axes[1].set_title(f"Rolling {ROLLING_WINDOW} Month Seasonal Coefficients")
    axes[1].legend(frameon=False, ncol=2, loc="upper right")
    axes[1].grid(axis="y")
    axes[1].xaxis.set_major_locator(mdates.YearLocator(1))
    axes[1].xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    fig.tight_layout()
    fig.savefig(rolling_path, bbox_inches="tight")
    plt.close(fig)
    return {"history": history_path, "fit": fit_path, "rolling": rolling_path}


def fmt_p(value: float) -> str:
    return "<0.001" if value < 0.001 else f"{value:.3f}"


def set_run_font(run, size=9.5, bold=False, color=None) -> None:
    run.font.name = "Arial"
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), "Arial")
    rpr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    if color is not None:
        run.font.color.rgb = color


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)
    normal = doc.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.2)
    normal.paragraph_format.space_after = Pt(5)
    for style_name, size in (("Title", 24), ("Heading 1", 16), ("Heading 2", 12.5)):
        style = doc.styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("Page ")
    set_run_font(run, size=8.5, color=RGBColor(107, 114, 128))
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, end])


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        available = 8.5 - 0.78 - 0.78
        scale = min(1.0, available / sum(widths))
        widths = [width * scale for width in widths]
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
    for index, text in enumerate(headers):
        cell = table.rows[0].cells[index]
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "24476B")
        cell._tc.get_or_add_tcPr().append(shading)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_run_font(paragraph.add_run(str(text)), size=8.5, bold=True, color=RGBColor(255, 255, 255))
    for row_index, values in enumerate(rows):
        row = table.add_row()
        for column_index, value in enumerate(values):
            cell = row.cells[column_index]
            if widths:
                cell.width = Inches(widths[column_index])
            if row_index % 2:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), "EEF4F8")
                cell._tc.get_or_add_tcPr().append(shading)
            paragraph = cell.paragraphs[0]
            paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT if column_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
            set_run_font(paragraph.add_run(str(value)), size=8.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def add_picture(doc: Document, path: Path, caption: str, width=6.55) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", caption)
    caption_paragraph = doc.add_paragraph()
    caption_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(caption_paragraph.add_run(caption), size=9, bold=True)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.keep_with_next = True
    set_run_font(paragraph.add_run(text), size=9.2, bold=True)


def equation_text(results: Results) -> str:
    aliases = {"y_lag1": "y(t-1)", "y_lag2": "y(t-2)", "January": "Jan(t)", "December": "Dec(t)"}
    parts = [f"y(t) = {results.model.params['const']:.6f}"]
    for term in results.selected_terms:
        parts.append(f"{results.model.params[term]:+.6f} {aliases.get(term, term)}")
    return " ".join(parts)


def build_report(results: Results, figures: dict[str, Path], output_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Consumer Card Fee Revenue Model Results"
    doc.add_heading("Consumer Card Fee Revenue Model Results", 0)
    subtitle = doc.add_paragraph()
    set_run_font(
        subtitle.add_run(
            f"In sample {IN_SAMPLE_START:%Y-%m} to {IN_SAMPLE_END:%Y-%m} and "
            f"out of time {OOT_START:%Y-%m} to {OOT_END:%Y-%m}"
        ),
        size=12,
        color=RGBColor(55, 65, 81),
    )

    diagnostics = model_diagnostics(results.model)
    accuracy = results.accuracy.set_index("window")
    macro_terms = [term for term in results.selected_terms if term not in BASE_TERMS]
    doc.add_heading("Executive Summary", level=1)
    doc.add_paragraph(
        "The model is estimated directly from the supplied workbooks. Forward selection retains "
        f"{', '.join(macro_terms)} in addition to two autoregressive terms and January and December indicators."
    )
    doc.add_paragraph(
        f"The in-sample regression uses {diagnostics['n']} observations, with R squared {diagnostics['r2']:.3f}, "
        f"adjusted R squared {diagnostics['adj_r2']:.3f}, Durbin Watson {diagnostics['dw']:.3f}, and Jarque Bera "
        f"p value {diagnostics['jb_p']:.3f}. The one-step OOT backtest has MAPE "
        f"{accuracy.loc['Out of time', 'mape_pct']:.2f} percent and RMSE {accuracy.loc['Out of time', 'rmse']:.2f}."
    )
    add_table(
        doc,
        ["Window", "n", "MAPE", "RMSE", "MAE"],
        [[row["window"], str(int(row["n"])), f"{row['mape_pct']:.3f}%", f"{row['rmse']:.3f}", f"{row['mae']:.3f}"] for _, row in results.accuracy.iterrows()],
        [1.7, 0.7, 1.2, 1.2, 1.2],
    )
    add_picture(doc, figures["history"], "Figure 1 Consumer Card fee revenue and the OOT period")

    doc.add_page_break()
    doc.add_heading("Data and Methodology", level=1)
    add_caption(doc, "Table 1 Macro candidate universe")
    add_table(
        doc,
        ["Group", "Variable", "Source field", "Transform"],
        [[spec["group"], name, spec["source"], spec["transform"]] for name, spec in MACRO_SPECS.items()],
        [0.9, 1.3, 2.8, 1.9],
    )
    doc.add_paragraph(
        "The dependent variable is 100 times the monthly log difference of Consumer Card fee revenue. Each macro "
        "series contributes its current transformation and a trailing three-month arithmetic average. Forward "
        "selection starts with y(t-1), y(t-2), January, and December and then applies the configured diagnostic gates."
    )
    first_round = results.screening.loc[results.screening["round"] == 1].sort_values("p_value")
    add_caption(doc, "Table 2 First-round macro screen")
    add_table(
        doc,
        ["Candidate", "Coefficient", "p value", "Adjusted R squared", "Pass"],
        [[row["candidate"], f"{row['coefficient']:.4f}", fmt_p(row["p_value"]), f"{row['adj_r2']:.4f}", "Yes" if row["passes"] else "No"] for _, row in first_round.iterrows()],
        [2.6, 1.0, 0.8, 1.5, 0.7],
    )
    add_caption(doc, "Table 3 Target stationarity tests")
    add_table(
        doc,
        ["Series", "n", "ADF p", "KPSS p"],
        [[row["series"], str(int(row["n"])), fmt_p(row["adf_p"]), fmt_p(row["kpss_p"])] for _, row in results.stationarity.iterrows()],
        [2.8, 0.8, 1.2, 1.2],
    )

    doc.add_page_break()
    doc.add_heading("Final Model", level=1)
    equation = doc.add_paragraph()
    equation.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(equation.add_run(equation_text(results)), size=9.2)
    add_caption(doc, "Table 4 Coefficient estimates")
    add_table(
        doc,
        ["Term", "Coefficient", "Standard error", "p value", "VIF"],
        [[row["term"], f"{row['coefficient']:.6f}", f"{row['std_error']:.6f}", fmt_p(row["p_value"]), "-" if pd.isna(row["vif"]) else f"{row['vif']:.3f}"] for _, row in results.coefficients.iterrows()],
        [2.0, 1.3, 1.3, 1.0, 0.8],
    )
    add_caption(doc, "Table 5 Model diagnostics")
    add_table(
        doc,
        ["Metric", "Value"],
        [
            ["Observations", str(diagnostics["n"])],
            ["R squared", f"{diagnostics['r2']:.6f}"],
            ["Adjusted R squared", f"{diagnostics['adj_r2']:.6f}"],
            ["Durbin Watson", f"{diagnostics['dw']:.6f}"],
            ["Jarque Bera p value", f"{diagnostics['jb_p']:.6f}"],
            ["Maximum VIF", f"{results.coefficients['vif'].max():.6f}"],
        ],
        [3.2, 2.0],
    )
    doc.add_paragraph(
        "Level forecasts equal the observed prior level multiplied by exp(predicted log change / 100). MAPE divides "
        "each absolute level error by the corresponding actual level. RMSE is the square root of mean squared level error."
    )

    doc.add_page_break()
    doc.add_heading("Predictive Accuracy", level=1)
    add_picture(doc, figures["fit"], "Figure 2 Actual and one-step predicted Consumer Card fee revenue")
    add_caption(doc, "Table 6 Transformed residual diagnostics")
    add_table(
        doc,
        ["Window", "n", "Mean", "Standard deviation", "Durbin Watson", "Jarque Bera p"],
        [[row["window"], str(int(row["n"])), f"{row['mean']:.6f}", f"{row['std']:.6f}", f"{row['dw']:.6f}", f"{row['jb_p']:.6f}"] for _, row in results.residual_diagnostics.iterrows()],
        [1.5, 0.6, 1.0, 1.4, 1.2, 1.3],
    )
    doc.add_paragraph(
        "The OOT calculation is a one-step conditional backtest using observed lagged target changes and the observed "
        "prior target level each month. It is not a recursive forecast from the final in-sample month."
    )

    doc.add_page_break()
    doc.add_heading("Structural Stability", level=1)
    add_caption(doc, "Table 7 CUSUM results")
    add_table(doc, ["Sample", "n", "Statistic", "p value"], [[row["sample"], str(int(row["n"])), f"{row['statistic']:.6f}", f"{row['p_value']:.6f}"] for _, row in results.cusum.iterrows()], [2.0, 0.8, 1.3, 1.3])
    add_caption(doc, "Table 8 Chow tests")
    add_table(doc, ["Break date", "n pre", "n post", "F statistic", "p value"], [[row["break_date"], str(int(row["n_pre"])), str(int(row["n_post"])), f"{row['f_stat']:.6f}", f"{row['p_value']:.6f}"] for _, row in results.chow.iterrows()], [1.5, 0.9, 0.9, 1.3, 1.3])
    add_caption(doc, f"Table 9 Rolling {ROLLING_WINDOW}-month coefficient summary")
    add_table(doc, ["Term", "Minimum", "Maximum", "Mean", "Same-sign windows", "Windows"], [[row["term"], f"{row['min']:.6f}", f"{row['max']:.6f}", f"{row['mean']:.6f}", str(int(row["same_sign_count"])), str(int(row["windows"]))] for _, row in results.rolling_summary.iterrows()], [1.8, 1.0, 1.0, 1.0, 1.3, 0.8])
    add_picture(doc, figures["rolling"], f"Figure 3 Rolling {ROLLING_WINDOW}-month coefficient estimates")

    doc.add_page_break()
    doc.add_heading("Out of Time Reestimation", level=1)
    add_caption(doc, "Table 10 OOT-only coefficient estimates")
    add_table(
        doc,
        ["Term", "Coefficient", "Standard error", "p value", "VIF"],
        [[row["term"], f"{row['coefficient']:.6f}", f"{row['std_error']:.6f}", fmt_p(row["p_value"]), "-" if pd.isna(row["vif"]) else f"{row['vif']:.3f}"] for _, row in results.oot_only.iterrows()],
        [2.0, 1.3, 1.3, 1.0, 0.8],
    )
    doc.add_paragraph("The OOT-only regression is a small-sample stability diagnostic. The fixed in-sample model is used for the OOT backtest.")
    doc.add_heading("Limitations", level=2)
    for text in (
        "The target workbook does not identify the currency or scale of the Consumer Card field.",
        "The macro workbook does not include release dates or real-time vintages.",
        "The OOT exercise is conditional on observed lagged target values.",
        "Forward selection evaluates related candidates without a multiplicity adjustment.",
    ):
        doc.add_paragraph(text, style="List Bullet")

    doc.add_heading("Source Provenance", level=2)
    audit = results.source_audit
    add_table(
        doc,
        ["Source", "Coverage", "Rows", "SHA 256"],
        [
            [audit["card_filename"], f"{audit['card']['start']} to {audit['card']['end']}", str(audit["card"]["rows"]), audit["card_sha256"]],
            [audit["macro_filename"], f"{audit['macro']['start']} to {audit['macro']['end']}", str(audit["macro"]["rows"]), audit["macro_sha256"]],
        ],
        [1.5, 1.5, 0.6, 3.3],
    )
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--card-file", type=Path, default=DEFAULT_CARD)
    parser.add_argument("--macro-file", type=Path, default=DEFAULT_MACRO)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT)
    parser.add_argument("--report-name", default="Consumer_Card_Model_Independent_Results.docx")
    parser.add_argument("--skip-report", action="store_true", help="Do not create the Word report")
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_dir = args.output_dir.resolve()
    results = run_analysis(args.card_file.resolve(), args.macro_file.resolve())
    save_outputs(results, output_dir)
    figures = make_plots(results, output_dir / "assets")
    if not args.skip_report:
        report_path = output_dir / args.report_name
        build_report(results, figures, report_path)
        print(f"Report: {report_path}")
    print(f"Selected terms: {results.selected_terms}")
    print(results.coefficients.to_string(index=False))
    print(results.accuracy.to_string(index=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
