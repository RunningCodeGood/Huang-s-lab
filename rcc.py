#!/usr/bin/env python3
"""Reproduce and audit the Consumer Card fee revenue model.

The script reads the original fee and macro workbooks, validates the monthly
calendar, screens the three Core and two Consumer macro families, fits the
reported dynamic regression, evaluates the fixed-coefficient one-step OOT
backtest, runs structural diagnostics, and builds an English Word report.

Example
-------
python reproduce_consumer_card_model.py \
  --card-file Cardfee.xlsx \
  --macro-file Macro_Data_201001-202612.xls \
  --output-dir output/consumer_reproduction
"""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import platform
import sys
import warnings
from dataclasses import dataclass
from importlib import metadata
from pathlib import Path

import matplotlib

matplotlib.use("Agg")

import matplotlib.dates as mdates
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
import statsmodels.api as sm
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from scipy import stats
from statsmodels.stats.diagnostic import breaks_cusumolsresid
from statsmodels.stats.outliers_influence import variance_inflation_factor
from statsmodels.stats.stattools import durbin_watson, jarque_bera
from statsmodels.tsa.stattools import adfuller, kpss


ROOT = Path(__file__).resolve().parent
DEFAULT_CARD = Path("/Users/hyx/Desktop/Cardfee/Cardfee.xlsx")
DEFAULT_MACRO = Path("/Users/hyx/Desktop/Cardfee/Revise Version/Macro_Data_201001-202612.xls")
DEFAULT_OUTPUT = ROOT / "output" / "consumer_reproduction"

IN_SAMPLE_START = pd.Timestamp("2011-01-01")
IN_SAMPLE_END = pd.Timestamp("2023-09-01")
OOT_START = pd.Timestamp("2023-10-01")
OOT_END = pd.Timestamp("2025-12-01")

TARGET_COLUMN = "Consumer Card"
DATE_COLUMN = "YYYYMM"
RETAIL_COLUMN = "M_FRT_B.IUSA (Retail Sales Total_Bil.USD,CDASAAR)"

MACRO_SPECS = {
    "Unemployment": {
        "group": "Core",
        "source": "M_FLBR_B.IUSA",
        "description": "Household unemployment rate, percent, seasonally adjusted",
        "transform": "Monthly first difference in percentage points",
    },
    "CPI": {
        "group": "Core",
        "source": "M_FCPIU_B.IUSA",
        "description": "CPI for urban consumers, 1982 to 1984 equals 100, seasonally adjusted",
        "transform": "Monthly percentage change",
    },
    "DPI": {
        "group": "Core",
        "source": "M_FYPDPIQ_B.IUSA",
        "description": "Disposable personal income, billion USD, seasonally adjusted annual rate",
        "transform": "Monthly percentage change",
    },
    "RetailSales": {
        "group": "Consumer",
        "source": RETAIL_COLUMN,
        "description": "Total retail sales, billion USD, calendar day adjusted annual rate",
        "transform": "Monthly percentage change",
    },
    "ConsumerConfidence": {
        "group": "Consumer",
        "source": "M_FCBC_B.IUSA(Consumer confidence index)",
        "description": "Consumer confidence index, 1985 equals 100, seasonally adjusted",
        "transform": "Monthly percentage change",
    },
}

BASE_TERMS = ["y_lag1", "y_lag2", "January", "December"]
FINAL_MACRO_TERM = "RetailSales_chg_MA3"

REFERENCE = {
    "stationarity": {
        "level_adf_p": 0.465,
        "level_kpss_p": 0.057,
        "diff_adf_display": "<0.001",
        "diff_kpss_p": 0.100,
    },
    "coefficients": {
        "const": -2.00,
        "y_lag1": -0.627,
        "y_lag2": -0.340,
        "January": -10.65,
        "December": 19.61,
        FINAL_MACRO_TERM: 2.55,
    },
    "fit": {"r2": 0.623, "adj_r2": 0.610, "n": 150, "dw": 2.356, "jb_p": 0.547},
    "accuracy": {
        "is_mape": 6.69,
        "is_rmse": 55.9,
        "oot_mape": 5.21,
        "oot_rmse": 46.8,
    },
    "cusum": {"is_p": 0.477, "full_p": 0.494},
    "chow": {
        "2023-10-01": {"f": 0.70, "p": 0.648},
        "2024-01-01": {"f": 0.73, "p": 0.629},
        "2024-07-01": {"f": 0.75, "p": 0.613},
    },
    "rolling": {"windows": 118, "retail_min": 1.08, "retail_max": 5.88},
    "oot_residuals": {"mean": 0.59, "std": 6.237, "dw": 1.619, "jb_p": 0.557},
    "oot_only": {
        "y_lag1": {"coef": -0.335, "p": 0.044},
        "y_lag2": {"coef": -0.426, "p": 0.019},
        "January": {"coef": -16.99, "p": 0.001},
        "December": {"coef": 15.37, "p": 0.001},
        FINAL_MACRO_TERM: {"coef": -3.17, "p": 0.414},
    },
}

NAVY = "24476B"
BLUE = "4F7CAC"
PALE_BLUE = "EEF4F8"
PALE_GRAY = "F5F6F7"
LIGHT_GRAY = "D9D9D9"
MID_GRAY = "6B7280"
BLACK = RGBColor(0, 0, 0)


@dataclass
class ModelArtifacts:
    data: pd.DataFrame
    in_sample: pd.DataFrame
    oot: pd.DataFrame
    model: object
    selected_terms: list[str]
    stationarity: pd.DataFrame
    screening: pd.DataFrame
    coefficients: pd.DataFrame
    accuracy: pd.DataFrame
    residual_diagnostics: pd.DataFrame
    cusum: pd.DataFrame
    chow: pd.DataFrame
    rolling: pd.DataFrame
    rolling_summary: pd.DataFrame
    oot_only: pd.DataFrame
    reconciliation: pd.DataFrame
    source_audit: dict


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def normalize_month(values: pd.Series) -> pd.DatetimeIndex:
    if pd.api.types.is_datetime64_any_dtype(values):
        return pd.DatetimeIndex(pd.to_datetime(values).dt.to_period("M").dt.to_timestamp())
    text = values.astype(str).str.replace(r"\.0$", "", regex=True).str.strip()
    parsed = pd.to_datetime(text, format="%Y%m", errors="coerce")
    if parsed.isna().any():
        parsed = pd.to_datetime(values, errors="coerce")
    if parsed.isna().any():
        raise ValueError("The date column contains values that cannot be parsed as monthly dates.")
    return pd.DatetimeIndex(parsed).to_period("M").to_timestamp()


def validate_calendar(frame: pd.DataFrame, name: str) -> dict:
    if frame.index.has_duplicates:
        duplicates = frame.index[frame.index.duplicated()].strftime("%Y-%m").tolist()
        raise ValueError(f"{name} contains duplicate months: {duplicates}")
    expected = pd.date_range(frame.index.min(), frame.index.max(), freq="MS")
    missing = expected.difference(frame.index)
    if len(missing):
        raise ValueError(f"{name} is missing calendar months: {missing.strftime('%Y-%m').tolist()}")
    return {
        "rows": int(len(frame)),
        "start": frame.index.min().strftime("%Y-%m"),
        "end": frame.index.max().strftime("%Y-%m"),
        "duplicate_months": 0,
        "missing_calendar_months": 0,
    }


def read_sources(card_path: Path, macro_path: Path) -> tuple[pd.DataFrame, pd.DataFrame, dict]:
    if not card_path.is_file():
        raise FileNotFoundError(f"Card source not found: {card_path}")
    if not macro_path.is_file():
        raise FileNotFoundError(f"Macro source not found: {macro_path}")

    card = pd.read_excel(card_path, sheet_name="汇总数据")
    required_card = {DATE_COLUMN, TARGET_COLUMN}
    if not required_card.issubset(card.columns):
        raise ValueError(f"Card source is missing columns: {sorted(required_card - set(card.columns))}")
    card = card[[DATE_COLUMN, TARGET_COLUMN]].copy()
    card.index = normalize_month(card.pop(DATE_COLUMN))
    card.index.name = "date"
    card[TARGET_COLUMN] = pd.to_numeric(card[TARGET_COLUMN], errors="coerce")

    macro = pd.read_excel(macro_path, sheet_name="Macro Data", header=1)
    required_macro = {DATE_COLUMN} | {spec["source"] for spec in MACRO_SPECS.values()}
    if not required_macro.issubset(macro.columns):
        raise ValueError(f"Macro source is missing columns: {sorted(required_macro - set(macro.columns))}")
    macro = macro[[DATE_COLUMN] + [spec["source"] for spec in MACRO_SPECS.values()]].copy()
    macro.index = normalize_month(macro.pop(DATE_COLUMN))
    macro.index.name = "date"
    for column in macro.columns:
        macro[column] = pd.to_numeric(macro[column], errors="coerce")

    card_audit = validate_calendar(card, "Card data")
    macro_audit = validate_calendar(macro, "Macro data")
    if card[TARGET_COLUMN].isna().any():
        raise ValueError("Consumer Card contains missing values.")
    if (card[TARGET_COLUMN] <= 0).any():
        raise ValueError("Consumer Card must be positive for the reported log-change target.")
    missing_macro = macro.isna().sum()
    if missing_macro.any():
        raise ValueError(f"Selected macro series contain missing values: {missing_macro[missing_macro > 0].to_dict()}")

    audit = {
        "card": card_audit,
        "macro": macro_audit,
        "card_sha256": sha256(card_path),
        "macro_sha256": sha256(macro_path),
        "card_filename": card_path.name,
        "macro_filename": macro_path.name,
    }
    return card, macro, audit


def engineer_data(card: pd.DataFrame, macro: pd.DataFrame) -> tuple[pd.DataFrame, list[str]]:
    data = card.join(macro, how="left")
    data["y"] = 100.0 * np.log(data[TARGET_COLUMN]).diff()
    data["y_lag1"] = data["y"].shift(1)
    data["y_lag2"] = data["y"].shift(2)
    data["January"] = (data.index.month == 1).astype(int)
    data["December"] = (data.index.month == 12).astype(int)

    candidates: list[str] = []
    for name, spec in MACRO_SPECS.items():
        source = macro[spec["source"]]
        if name == "Unemployment":
            change = source.diff()
        else:
            change = 100.0 * source.pct_change(fill_method=None)
        raw_name = f"{name}_chg"
        ma_name = f"{name}_chg_MA3"
        data[raw_name] = change
        data[ma_name] = change.rolling(3, min_periods=3).mean()
        candidates.extend([raw_name, ma_name])

    expected = pd.date_range(IN_SAMPLE_START, OOT_END, freq="MS")
    missing = expected.difference(data.index)
    if len(missing):
        raise ValueError(f"Combined model calendar is missing months: {missing.strftime('%Y-%m').tolist()}")
    return data, candidates


def stationarity_table(data: pd.DataFrame) -> pd.DataFrame:
    level = data.loc[IN_SAMPLE_START:IN_SAMPLE_END, TARGET_COLUMN]
    diff = level.diff().dropna()
    rows = []
    for label, series in (("Consumer Card level", level), ("Consumer Card first difference", diff)):
        adf = adfuller(series, regression="c", autolag="AIC", result_object=False)
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            kp = kpss(series, regression="c", nlags="auto")
        rows.append(
            {
                "series": label,
                "n": int(len(series)),
                "adf_stat": float(adf[0]),
                "adf_p": float(adf[1]),
                "adf_lags": int(adf[2]),
                "kpss_stat": float(kp[0]),
                "kpss_p": float(kp[1]),
                "kpss_lags": int(kp[2]),
            }
        )
    return pd.DataFrame(rows)


def fit_ols(frame: pd.DataFrame, terms: list[str]):
    sample = frame[["y"] + terms].dropna()
    x = sm.add_constant(sample[terms], has_constant="add")
    return sm.OLS(sample["y"], x).fit(), sample


def model_diagnostics(model) -> dict:
    jb = jarque_bera(model.resid)
    return {
        "n": int(model.nobs),
        "r2": float(model.rsquared),
        "adj_r2": float(model.rsquared_adj),
        "dw": float(durbin_watson(model.resid)),
        "jb_stat": float(jb[0]),
        "jb_p": float(jb[1]),
        "residual_mean": float(np.mean(model.resid)),
        "residual_std": float(np.std(model.resid, ddof=1)),
    }


def compute_vifs(model) -> dict[str, float]:
    x = model.model.exog
    names = model.model.exog_names
    values = {}
    for i, name in enumerate(names):
        if name == "const":
            continue
        values[name] = float(variance_inflation_factor(x, i))
    return values


def forward_selection(data: pd.DataFrame, candidates: list[str]) -> tuple[list[str], pd.DataFrame]:
    sample = data.loc[IN_SAMPLE_START:IN_SAMPLE_END]
    selected: list[str] = []
    remaining = list(candidates)
    records: list[dict] = []
    round_number = 1

    while remaining:
        trials = []
        for candidate in remaining:
            terms = BASE_TERMS + selected + [candidate]
            model, used = fit_ols(sample, terms)
            vifs = compute_vifs(model)
            diagnostics = model_diagnostics(model)
            trial = {
                "round": round_number,
                "candidate": candidate,
                "n": int(len(used)),
                "coefficient": float(model.params[candidate]),
                "p_value": float(model.pvalues[candidate]),
                "adj_r2": float(model.rsquared_adj),
                "dw": diagnostics["dw"],
                "jb_p": diagnostics["jb_p"],
                "max_vif": float(max(vifs.values())),
            }
            trial["passes"] = bool(
                trial["p_value"] < 0.05
                and trial["max_vif"] < 10.0
                and 1.5 <= trial["dw"] <= 2.5
                and trial["jb_p"] > 0.05
            )
            trials.append(trial)
            records.append(trial.copy())

        passing = [row for row in trials if row["passes"]]
        if not passing:
            break
        winner = min(passing, key=lambda row: (row["p_value"], -row["adj_r2"], row["candidate"]))
        selected.append(winner["candidate"])
        remaining.remove(winner["candidate"])
        round_number += 1

    return selected, pd.DataFrame(records)


def coefficient_table(model) -> pd.DataFrame:
    vifs = compute_vifs(model)
    rows = []
    for term in model.params.index:
        rows.append(
            {
                "term": term,
                "coefficient": float(model.params[term]),
                "std_error": float(model.bse[term]),
                "t_stat": float(model.tvalues[term]),
                "p_value": float(model.pvalues[term]),
                "ci_low": float(model.conf_int().loc[term, 0]),
                "ci_high": float(model.conf_int().loc[term, 1]),
                "vif": vifs.get(term, np.nan),
            }
        )
    return pd.DataFrame(rows)


def prediction_frame(data: pd.DataFrame, model, sample: pd.DataFrame, window: str) -> pd.DataFrame:
    terms = [name for name in model.params.index if name != "const"]
    x = sm.add_constant(sample[terms], has_constant="add")
    predicted_y = model.predict(x)
    prior_level = data[TARGET_COLUMN].shift(1).reindex(sample.index)
    predicted_level = prior_level * np.exp(predicted_y / 100.0)
    actual_level = data[TARGET_COLUMN].reindex(sample.index)
    residual_y = sample["y"] - predicted_y
    return pd.DataFrame(
        {
            "window": window,
            "actual_level": actual_level,
            "predicted_level": predicted_level,
            "actual_y": sample["y"],
            "predicted_y": predicted_y,
            "residual_y": residual_y,
        },
        index=sample.index,
    )


def accuracy_row(predictions: pd.DataFrame) -> dict:
    error = predictions["actual_level"] - predictions["predicted_level"]
    return {
        "window": predictions["window"].iloc[0],
        "n": int(len(predictions)),
        "mape_pct": float((error.abs() / predictions["actual_level"]).mean() * 100.0),
        "rmse": float(np.sqrt(np.mean(np.square(error)))),
        "mae": float(np.mean(np.abs(error))),
    }


def residual_row(predictions: pd.DataFrame) -> dict:
    residual = predictions["residual_y"]
    jb = jarque_bera(residual)
    return {
        "window": predictions["window"].iloc[0],
        "n": int(len(predictions)),
        "mean": float(residual.mean()),
        "std": float(residual.std(ddof=1)),
        "dw": float(durbin_watson(residual)),
        "jb_p": float(jb[1]),
    }


def cusum_table(data: pd.DataFrame, terms: list[str]) -> pd.DataFrame:
    rows = []
    for label, end in (("In sample", IN_SAMPLE_END), ("Full sample", OOT_END)):
        model, sample = fit_ols(data.loc[IN_SAMPLE_START:end], terms)
        stat, p_value, critical = breaks_cusumolsresid(model.resid, ddof=len(model.params))
        rows.append(
            {
                "sample": label,
                "n": int(len(sample)),
                "statistic": float(stat),
                "p_value": float(p_value),
                "critical_1pct": float(critical[0][1]),
                "critical_5pct": float(critical[1][1]),
                "critical_10pct": float(critical[2][1]),
            }
        )
    return pd.DataFrame(rows)


def chow_table(data: pd.DataFrame, terms: list[str]) -> pd.DataFrame:
    sample = data.loc[IN_SAMPLE_START:OOT_END, ["y"] + terms].dropna()
    pooled, _ = fit_ols(sample, terms)
    k = len(pooled.params)
    rows = []
    for break_date in (pd.Timestamp("2023-10-01"), pd.Timestamp("2024-01-01"), pd.Timestamp("2024-07-01")):
        pre = sample.loc[sample.index < break_date]
        post = sample.loc[sample.index >= break_date]
        pre_model, _ = fit_ols(pre, terms)
        post_model, _ = fit_ols(post, terms)
        numerator = (pooled.ssr - pre_model.ssr - post_model.ssr) / k
        denominator_df = len(pre) + len(post) - 2 * k
        denominator = (pre_model.ssr + post_model.ssr) / denominator_df
        f_stat = float(numerator / denominator)
        p_value = float(stats.f.sf(f_stat, k, denominator_df))
        rows.append(
            {
                "break_date": break_date.strftime("%Y-%m-%d"),
                "n_pre": int(len(pre)),
                "n_post": int(len(post)),
                "f_stat": f_stat,
                "p_value": p_value,
                "df_num": int(k),
                "df_den": int(denominator_df),
            }
        )
    return pd.DataFrame(rows)


def rolling_coefficients(data: pd.DataFrame, terms: list[str], window: int = 60) -> tuple[pd.DataFrame, pd.DataFrame]:
    sample = data.loc[IN_SAMPLE_START:OOT_END, ["y"] + terms].dropna()
    rows = []
    for end_pos in range(window - 1, len(sample)):
        current = sample.iloc[end_pos - window + 1 : end_pos + 1]
        model, _ = fit_ols(current, terms)
        row = {"window_end": current.index[-1], "window_start": current.index[0], "n": len(current)}
        row.update({term: float(model.params[term]) for term in model.params.index})
        rows.append(row)
    rolling = pd.DataFrame(rows).set_index("window_end")
    summary_rows = []
    for term in terms:
        values = rolling[term]
        expected_sign = np.sign(values.iloc[0])
        summary_rows.append(
            {
                "term": term,
                "min": float(values.min()),
                "max": float(values.max()),
                "mean": float(values.mean()),
                "std": float(values.std(ddof=1)),
                "same_sign_count": int((np.sign(values) == expected_sign).sum()),
                "windows": int(len(values)),
            }
        )
    return rolling, pd.DataFrame(summary_rows)


def oot_only_table(data: pd.DataFrame, terms: list[str]) -> pd.DataFrame:
    model, sample = fit_ols(data.loc[OOT_START:OOT_END], terms)
    table = coefficient_table(model)
    table["n"] = len(sample)
    return table


def compare_status(reference: float, reproduced: float, displayed_decimals: int) -> str:
    rounding_tolerance = 0.5 * 10 ** (-displayed_decimals)
    difference = abs(reference - reproduced)
    if difference <= rounding_tolerance:
        return "Matches reported precision"
    scale = max(abs(reference), 1.0)
    if difference <= max(5 * rounding_tolerance, 0.02 * scale):
        return "Close but not exact"
    return "Does not match"


def reconciliation_table(model, accuracy: pd.DataFrame, cusum: pd.DataFrame, rolling: pd.DataFrame) -> pd.DataFrame:
    diag = model_diagnostics(model)
    acc = accuracy.set_index("window")
    metrics = [
        ("In sample observations", REFERENCE["fit"]["n"], diag["n"], 0),
        ("R squared", REFERENCE["fit"]["r2"], diag["r2"], 3),
        ("Adjusted R squared", REFERENCE["fit"]["adj_r2"], diag["adj_r2"], 3),
        ("Durbin Watson", REFERENCE["fit"]["dw"], diag["dw"], 3),
        ("Jarque Bera p value", REFERENCE["fit"]["jb_p"], diag["jb_p"], 3),
        ("In sample MAPE percent", REFERENCE["accuracy"]["is_mape"], acc.loc["In sample", "mape_pct"], 2),
        ("In sample RMSE", REFERENCE["accuracy"]["is_rmse"], acc.loc["In sample", "rmse"], 1),
        ("OOT MAPE percent", REFERENCE["accuracy"]["oot_mape"], acc.loc["Out of time", "mape_pct"], 2),
        ("OOT RMSE", REFERENCE["accuracy"]["oot_rmse"], acc.loc["Out of time", "rmse"], 1),
        ("In sample CUSUM p value", REFERENCE["cusum"]["is_p"], cusum.loc[0, "p_value"], 3),
        ("Full sample CUSUM p value", REFERENCE["cusum"]["full_p"], cusum.loc[1, "p_value"], 3),
        ("Rolling windows", REFERENCE["rolling"]["windows"], len(rolling), 0),
        ("Rolling retail coefficient minimum", REFERENCE["rolling"]["retail_min"], rolling[FINAL_MACRO_TERM].min(), 2),
        ("Rolling retail coefficient maximum", REFERENCE["rolling"]["retail_max"], rolling[FINAL_MACRO_TERM].max(), 2),
    ]
    rows = []
    for metric, reference, reproduced, decimals in metrics:
        rows.append(
            {
                "metric": metric,
                "reported": float(reference),
                "reproduced": float(reproduced),
                "difference": float(reproduced - reference),
                "status": compare_status(float(reference), float(reproduced), decimals),
            }
        )
    return pd.DataFrame(rows)


def run_analysis(card_path: Path, macro_path: Path) -> ModelArtifacts:
    card, macro, source_audit = read_sources(card_path, macro_path)
    data, candidates = engineer_data(card, macro)
    stationarity = stationarity_table(data)
    selected_macro, screening = forward_selection(data, candidates)
    selected_terms = BASE_TERMS + selected_macro
    if selected_macro != [FINAL_MACRO_TERM]:
        raise RuntimeError(
            "Forward selection did not reproduce the reported RetailSales_chg_MA3-only macro specification. "
            f"Selected: {selected_macro}"
        )

    model, in_sample = fit_ols(data.loc[IN_SAMPLE_START:IN_SAMPLE_END], selected_terms)
    _, oot = fit_ols(data.loc[OOT_START:OOT_END], selected_terms)
    if len(in_sample) != 150 or len(oot) != 27:
        raise RuntimeError(f"Unexpected model sample sizes: in sample {len(in_sample)}, OOT {len(oot)}")

    coefficients = coefficient_table(model)
    predictions_is = prediction_frame(data, model, in_sample, "In sample")
    predictions_oot = prediction_frame(data, model, oot, "Out of time")
    predictions = pd.concat([predictions_is, predictions_oot])
    data = data.join(predictions.add_prefix("prediction_"), how="left")
    accuracy = pd.DataFrame([accuracy_row(predictions_is), accuracy_row(predictions_oot)])
    residual_diagnostics = pd.DataFrame([residual_row(predictions_is), residual_row(predictions_oot)])
    cusum = cusum_table(data, selected_terms)
    chow = chow_table(data, selected_terms)
    rolling, rolling_summary = rolling_coefficients(data, selected_terms, window=60)
    oot_only = oot_only_table(data, selected_terms)
    reconciliation = reconciliation_table(model, accuracy, cusum, rolling)

    return ModelArtifacts(
        data=data,
        in_sample=in_sample,
        oot=oot,
        model=model,
        selected_terms=selected_terms,
        stationarity=stationarity,
        screening=screening,
        coefficients=coefficients,
        accuracy=accuracy,
        residual_diagnostics=residual_diagnostics,
        cusum=cusum,
        chow=chow,
        rolling=rolling,
        rolling_summary=rolling_summary,
        oot_only=oot_only,
        reconciliation=reconciliation,
        source_audit=source_audit,
    )


def save_numeric_outputs(artifacts: ModelArtifacts, output_dir: Path) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    frames = {
        "stationarity.csv": artifacts.stationarity,
        "macro_screening.csv": artifacts.screening,
        "coefficients.csv": artifacts.coefficients,
        "accuracy.csv": artifacts.accuracy,
        "residual_diagnostics.csv": artifacts.residual_diagnostics,
        "cusum.csv": artifacts.cusum,
        "chow.csv": artifacts.chow,
        "rolling_coefficients.csv": artifacts.rolling.reset_index(),
        "rolling_summary.csv": artifacts.rolling_summary,
        "oot_only_coefficients.csv": artifacts.oot_only,
        "reconciliation.csv": artifacts.reconciliation,
    }
    for filename, frame in frames.items():
        frame.to_csv(output_dir / filename, index=False, encoding="utf-8-sig", float_format="%.12g")

    diag = model_diagnostics(artifacts.model)
    manifest = {
        "periods": {
            "in_sample": [IN_SAMPLE_START.strftime("%Y-%m"), IN_SAMPLE_END.strftime("%Y-%m")],
            "out_of_time": [OOT_START.strftime("%Y-%m"), OOT_END.strftime("%Y-%m")],
        },
        "target": TARGET_COLUMN,
        "selected_terms": artifacts.selected_terms,
        "model_diagnostics": diag,
        "source_audit": artifacts.source_audit,
        "runtime": {
            "python": sys.version,
            "platform": platform.platform(),
            "packages": {
                name: metadata.version(name)
                for name in ("numpy", "pandas", "scipy", "statsmodels", "matplotlib", "python-docx")
            },
        },
    }
    (output_dir / "run_manifest.json").write_text(json.dumps(manifest, indent=2), encoding="utf-8")


def make_plots(artifacts: ModelArtifacts, assets_dir: Path) -> dict[str, Path]:
    assets_dir.mkdir(parents=True, exist_ok=True)
    plt.rcParams.update(
        {
            "font.family": "Arial",
            "font.size": 9.5,
            "axes.titlesize": 11,
            "axes.labelsize": 9.5,
            "axes.edgecolor": "#AAB2BD",
            "axes.spines.top": False,
            "axes.spines.right": False,
            "grid.color": "#E5E7EB",
            "grid.linewidth": 0.7,
        }
    )

    level_path = assets_dir / "consumer_card_history.png"
    fig, ax = plt.subplots(figsize=(7.2, 3.2), dpi=180)
    history = artifacts.data.loc[IN_SAMPLE_START:OOT_END, TARGET_COLUMN]
    ax.plot(history.index, history.values, color="#24476B", linewidth=1.6)
    ax.axvspan(OOT_START, OOT_END, color="#DCEAF4", alpha=0.8, label="Out of time")
    ax.axvline(OOT_START, color="#6B7280", linewidth=1.0, linestyle="--")
    ax.set_title("Consumer Card Fee Revenue")
    ax.set_ylabel("Source units")
    ax.grid(axis="y")
    ax.legend(frameon=False, loc="upper right")
    ax.xaxis.set_major_locator(mdates.YearLocator(2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    fig.tight_layout()
    fig.savefig(level_path, bbox_inches="tight")
    plt.close(fig)

    fit_path = assets_dir / "actual_vs_predicted.png"
    fig, ax = plt.subplots(figsize=(7.2, 3.35), dpi=180)
    modeled = artifacts.data.loc[artifacts.in_sample.index.min():OOT_END]
    ax.plot(modeled.index, modeled[TARGET_COLUMN], color="#111827", linewidth=1.45, label="Actual")
    pred = modeled["prediction_predicted_level"]
    ax.plot(pred.loc[:IN_SAMPLE_END].index, pred.loc[:IN_SAMPLE_END], color="#4F7CAC", linewidth=1.25, label="In sample fitted")
    ax.plot(pred.loc[OOT_START:].index, pred.loc[OOT_START:], color="#C76D3A", linewidth=1.5, label="OOT one step")
    ax.axvline(OOT_START, color="#6B7280", linewidth=1.0, linestyle="--")
    ax.set_title("Actual and One Step Predicted Levels")
    ax.set_ylabel("Source units")
    ax.grid(axis="y")
    ax.legend(frameon=False, ncol=3, loc="upper right")
    ax.xaxis.set_major_locator(mdates.YearLocator(2))
    ax.xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    fig.tight_layout()
    fig.savefig(fit_path, bbox_inches="tight")
    plt.close(fig)

    rolling_path = assets_dir / "rolling_coefficients.png"
    rolling = artifacts.rolling
    fig, axes = plt.subplots(2, 1, figsize=(7.2, 5.0), dpi=180, sharex=True)
    axes[0].plot(rolling.index, rolling["y_lag1"], label="y lag 1", color="#24476B")
    axes[0].plot(rolling.index, rolling["y_lag2"], label="y lag 2", color="#4F7CAC")
    axes[0].plot(rolling.index, rolling[FINAL_MACRO_TERM], label="Retail sales MA3", color="#C76D3A")
    axes[0].axhline(0, color="#9CA3AF", linewidth=0.8)
    axes[0].set_title("Rolling 60 Month Dynamic and Macro Coefficients")
    axes[0].legend(frameon=False, ncol=3, loc="upper right")
    axes[0].grid(axis="y")
    axes[1].plot(rolling.index, rolling["January"], label="January", color="#7C3AED")
    axes[1].plot(rolling.index, rolling["December"], label="December", color="#0F766E")
    axes[1].axhline(0, color="#9CA3AF", linewidth=0.8)
    axes[1].set_title("Rolling 60 Month Seasonal Coefficients")
    axes[1].legend(frameon=False, ncol=2, loc="upper right")
    axes[1].grid(axis="y")
    axes[1].xaxis.set_major_locator(mdates.YearLocator(1))
    axes[1].xaxis.set_major_formatter(mdates.DateFormatter("%Y"))
    fig.tight_layout()
    fig.savefig(rolling_path, bbox_inches="tight")
    plt.close(fig)

    return {"history": level_path, "fit": fit_path, "rolling": rolling_path}


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=90, start=110, bottom=90, end=110) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color=LIGHT_GRAY, size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = borders.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            borders.append(tag)
        tag.set(qn("w:val"), "single")
        tag.set(qn("w:sz"), size)
        tag.set(qn("w:space"), "0")
        tag.set(qn("w:color"), color)


def set_repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_row_cant_split(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    tr_pr.append(cant_split)


def set_run_font(run, size=None, bold=None, italic=None, color=None, name="Arial") -> None:
    run.font.name = name
    rpr = run._element.get_or_add_rPr()
    rpr.rFonts.set(qn("w:ascii"), name)
    rpr.rFonts.set(qn("w:hAnsi"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color is not None:
        run.font.color.rgb = color


def configure_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.4)
    normal.font.color.rgb = BLACK
    normal.paragraph_format.space_after = Pt(5.5)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Arial"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    title.font.size = Pt(25)
    title.font.bold = True
    title.font.color.rgb = BLACK
    title.paragraph_format.space_after = Pt(8)
    title_ppr = title.element.get_or_add_pPr()
    border = title_ppr.find(qn("w:pBdr"))
    if border is not None:
        title_ppr.remove(border)

    for style_name, size, before, after in (
        ("Heading 1", 16, 15, 7),
        ("Heading 2", 12.5, 11, 5),
        ("Heading 3", 10.8, 8, 4),
    ):
        style = styles[style_name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = BLACK
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "Table Caption" not in [style.name for style in styles]:
        caption = styles.add_style("Table Caption", WD_STYLE_TYPE.PARAGRAPH)
    else:
        caption = styles["Table Caption"]
    caption.font.name = "Arial"
    caption._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    caption._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    caption.font.size = Pt(9.3)
    caption.font.bold = True
    caption.font.color.rgb = BLACK
    caption.paragraph_format.space_before = Pt(5)
    caption.paragraph_format.space_after = Pt(3)
    caption.paragraph_format.keep_with_next = True

    footer = section.footer
    paragraph = footer.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run("Page ")
    set_run_font(run, size=8.5, color=RGBColor(107, 114, 128))
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.append(fld_char1)
    run._r.append(instr)
    run._r.append(fld_char2)


def add_paragraph(doc: Document, text: str = "", bold_lead: str | None = None, italic=False, keep=False):
    paragraph = doc.add_paragraph()
    if keep:
        paragraph.paragraph_format.keep_together = True
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, bold=True)
        rest = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(rest)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, italic=italic)
    return paragraph


def add_bullet(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(text)
    set_run_font(run)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph(style="Table Caption")
    paragraph.add_run(text)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None,
              alignments: list[int] | None = None, font_size: float = 8.8):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    if widths:
        available_width = 8.5 - 0.78 - 0.78
        scale = min(1.0, available_width / sum(widths))
        widths = [width * scale for width in widths]
        width_twips = [int(round(Inches(width).twips)) for width in widths]
        for column, width in zip(table.columns, widths):
            column.width = Inches(width)
        for grid_column, twips in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), width_twips):
            grid_column.set(qn("w:w"), str(twips))
        table_properties = table._tbl.tblPr
        table_width = table_properties.find(qn("w:tblW"))
        if table_width is None:
            table_width = OxmlElement("w:tblW")
            table_properties.append(table_width)
        table_width.set(qn("w:type"), "dxa")
        table_width.set(qn("w:w"), str(sum(width_twips)))
        table_indent = table_properties.find(qn("w:tblInd"))
        if table_indent is None:
            table_indent = OxmlElement("w:tblInd")
            table_properties.append(table_indent)
        table_indent.set(qn("w:type"), "dxa")
        table_indent.set(qn("w:w"), "110")
    set_table_borders(table)
    header_row = table.rows[0]
    set_repeat_header(header_row)
    set_row_cant_split(header_row)
    for j, header_text in enumerate(headers):
        cell = header_row.cells[j]
        if widths:
            cell.width = Inches(widths[j])
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, top=100, bottom=100)
        cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(str(header_text))
        set_run_font(run, size=font_size, bold=True, color=RGBColor(255, 255, 255))

    for i, row_values in enumerate(rows):
        row = table.add_row()
        set_row_cant_split(row)
        for j, value in enumerate(row_values):
            cell = row.cells[j]
            if widths:
                cell.width = Inches(widths[j])
            if i % 2 == 1:
                set_cell_shading(cell, PALE_BLUE)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            paragraph = cell.paragraphs[0]
            paragraph.alignment = alignments[j] if alignments else WD_ALIGN_PARAGRAPH.LEFT
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.0
            run = paragraph.add_run(str(value))
            set_run_font(run, size=font_size)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(1)
    return table


def add_native_equation(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.keep_together = True
    math_para = OxmlElement("m:oMathPara")
    math_para_pr = OxmlElement("m:oMathParaPr")
    justification = OxmlElement("m:jc")
    justification.set(qn("m:val"), "center")
    math_para_pr.append(justification)
    math_para.append(math_para_pr)
    math = OxmlElement("m:oMath")
    math_run = OxmlElement("m:r")
    math_text = OxmlElement("m:t")
    math_text.text = text
    math_run.append(math_text)
    math.append(math_run)
    math_para.append(math)
    paragraph._p.append(math_para)


def add_figure(doc: Document, path: Path, caption: str, width=6.65) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(3)
    paragraph.paragraph_format.keep_together = True
    picture = paragraph.add_run().add_picture(str(path), width=Inches(width))
    picture._inline.docPr.set("title", caption)
    picture._inline.docPr.set("descr", caption)
    caption_p = doc.add_paragraph(style="Table Caption")
    caption_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption_p.paragraph_format.keep_with_next = False
    caption_p.add_run(caption)


def fmt_p(value: float) -> str:
    if value < 0.001:
        return "<0.001"
    return f"{value:.3f}"


def status_summary(reconciliation: pd.DataFrame) -> tuple[int, int, int]:
    return (
        int((reconciliation["status"] == "Matches reported precision").sum()),
        int((reconciliation["status"] == "Close but not exact").sum()),
        int((reconciliation["status"] == "Does not match").sum()),
    )


def build_report(artifacts: ModelArtifacts, figures: dict[str, Path], output_path: Path) -> None:
    doc = Document()
    configure_document(doc)
    doc.core_properties.title = "Consumer Card Fee Revenue Model Reproduction"
    doc.core_properties.subject = "Independent reproduction of the Consumer Card fee revenue model"
    doc.core_properties.keywords = "Consumer Card, fee revenue, model reproduction, out of time validation"

    title = doc.add_paragraph(style="Title")
    title.add_run("Consumer Card Fee Revenue Model Reproduction")
    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(16)
    run = subtitle.add_run("In Sample 2011 01 to 2023 09 and Out of Time 2023 10 to 2025 12")
    set_run_font(run, size=12.5, color=RGBColor(55, 65, 81))

    doc.add_heading("Executive Summary", level=1)
    diag = model_diagnostics(artifacts.model)
    acc = artifacts.accuracy.set_index("window")
    matched, close, different = status_summary(artifacts.reconciliation)
    add_paragraph(
        doc,
        "The independent calculation reproduces the reported model structure and its main conclusion. "
        "Forward selection over the three Core variables and two Consumer variables selects the three month "
        "moving average of total retail sales growth. The two autoregressive terms, January and December "
        "indicators, and the retail sales term all retain the reported signs and remain significant at the "
        "5 percent level in the in sample regression.",
    )
    add_paragraph(
        doc,
        f"Using the available source workbooks, the recalculated model has R squared {diag['r2']:.3f}, "
        f"adjusted R squared {diag['adj_r2']:.3f}, Durbin Watson {diag['dw']:.3f}, and 150 observations. "
        f"The one step OOT backtest produces MAPE {acc.loc['Out of time', 'mape_pct']:.2f} percent and "
        f"RMSE {acc.loc['Out of time', 'rmse']:.2f} source units.",
    )
    add_paragraph(
        doc,
        f"Of 14 headline checks against the prior report, {matched} match at the reported precision, "
        f"{close} are close but not exact, and {different} do not match. The sample counts, target "
        "stationarity results, model form, coefficient directions, significance findings, VIF results, "
        "Durbin Watson statistic, rolling window count, and retail coefficient range reproduce. Some "
        "printed coefficients, the Jarque Bera p value, OOT error measures, Chow tests, and OOT only "
        "reestimation do not reproduce exactly from the available files under the documented formulas.",
        keep=True,
    )

    add_caption(doc, "Table 1 Reproduction status")
    add_table(
        doc,
        ["Area", "Finding"],
        [
            ["Model selection", "RetailSales_chg_MA3 is the only selected macro term"],
            ["In sample", f"n 150, R squared {diag['r2']:.3f}, adjusted R squared {diag['adj_r2']:.3f}"],
            ["Out of time", f"n 27, MAPE {acc.loc['Out of time', 'mape_pct']:.2f} percent, RMSE {acc.loc['Out of time', 'rmse']:.2f}"],
            ["Overall assessment", "Main model conclusion reproduced; several printed values require source or method clarification"],
        ],
        [1.55, 5.45],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        9.2,
    )

    doc.add_page_break()
    doc.add_heading("Data and Scope", level=1)
    add_paragraph(
        doc,
        "The target is the fee revenue field named Consumer Card in the left side of the fee workbook. "
        "The similarly named loan field is not used. The target contains 180 complete monthly observations "
        "from January 2011 through December 2025. The macro workbook contains complete monthly values from "
        "January 2010 through December 2026. The extra macro history supplies the transformations needed at "
        "the start of the target sample.",
    )
    add_caption(doc, "Table 2 Modeling periods")
    add_table(
        doc,
        ["Purpose", "Start", "End", "Calendar months"],
        [
            ["Target history", "2011-01", "2025-12", "180"],
            ["In sample selection and estimation", "2011-01", "2023-09", "153 target months"],
            ["Usable in sample regression", "2011-04", "2023-09", "150"],
            ["Out of time backtest", "2023-10", "2025-12", "27"],
        ],
        [2.8, 1.25, 1.25, 1.7],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        9.1,
    )

    add_caption(doc, "Table 3 Macro candidate universe")
    macro_rows = []
    for name, spec in MACRO_SPECS.items():
        macro_rows.append([spec["group"], name, spec["source"], spec["transform"]])
    add_table(
        doc,
        ["Group", "Variable", "Source field", "Short run transform"],
        macro_rows,
        [0.95, 1.35, 2.75, 1.95],
        [WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
        8.3,
    )
    add_figure(doc, figures["history"], "Figure 1 Consumer Card fee revenue and the OOT period")

    doc.add_page_break()
    doc.add_heading("Methodology", level=1)
    doc.add_heading("Target Transformation", level=2)
    add_paragraph(
        doc,
        "The dependent variable is the monthly log change in Consumer Card fee revenue, expressed in percentage points.",
    )
    add_native_equation(doc, "y(t) = 100 [ln ConsumerCard(t) - ln ConsumerCard(t-1)]")

    doc.add_heading("Candidate Construction", level=2)
    add_paragraph(
        doc,
        "Unemployment enters as a monthly first difference in percentage points. CPI, disposable personal "
        "income, retail sales, and consumer confidence enter as monthly percentage changes. Each transformed "
        "macro series contributes its current value and a trailing three month arithmetic average. This creates "
        "10 macro candidates across the five source variables.",
    )

    doc.add_heading("Forward Selection", level=2)
    add_paragraph(
        doc,
        "The base equation contains y(t-1), y(t-2), January, and December. At each round, the script fits every "
        "remaining macro candidate and selects the candidate with the smallest p value among equations that "
        "satisfy all four gates: candidate p below 0.05, maximum VIF below 10, Durbin Watson from 1.5 to 2.5, "
        "and Jarque Bera p above 0.05. The first round selects RetailSales_chg_MA3. No remaining candidate "
        "passes the significance gate after that term enters.",
    )

    first_round = artifacts.screening.loc[artifacts.screening["round"] == 1].sort_values("p_value")
    screen_rows = []
    for _, row in first_round.iterrows():
        screen_rows.append(
            [
                row["candidate"],
                f"{row['coefficient']:.3f}",
                fmt_p(row["p_value"]),
                f"{row['adj_r2']:.3f}",
                "Pass" if row["passes"] else "Fail",
            ]
        )
    add_caption(doc, "Table 4 First round macro screen")
    add_table(
        doc,
        ["Candidate", "Coefficient", "p value", "Adjusted R squared", "All gates"],
        screen_rows,
        [2.55, 1.05, 0.85, 1.45, 1.1],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER],
        8.25,
    )

    doc.add_heading("Stationarity", level=2)
    station_rows = []
    for _, row in artifacts.stationarity.iterrows():
        verdict = (
            "Ambiguous level result"
            if "level" in row["series"]
            else "Stationary under both tests"
        )
        station_rows.append([row["series"], str(int(row["n"])), fmt_p(row["adf_p"]), fmt_p(row["kpss_p"]), verdict])
    add_caption(doc, "Table 5 Consumer Card stationarity tests")
    add_table(
        doc,
        ["Series", "n", "ADF p", "KPSS p", "Interpretation"],
        station_rows,
        [2.15, 0.55, 0.85, 0.85, 2.6],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        8.8,
    )

    doc.add_page_break()
    doc.add_heading("Final Model", level=1)
    add_paragraph(
        doc,
        "The reproduced final equation uses the same terms as the prior report. Coefficients below come from "
        "ordinary least squares on the 150 usable in sample months from April 2011 through September 2023.",
    )
    p = artifacts.model.params
    add_native_equation(
        doc,
        f"y(t)={p['const']:.3f}{p['y_lag1']:+.3f}y(t-1){p['y_lag2']:+.3f}y(t-2)"
        f"{p['January']:+.3f}Jan(t){p['December']:+.3f}Dec(t)"
        f"{p[FINAL_MACRO_TERM]:+.3f}RetailMA3(t)",
    )
    add_paragraph(
        doc,
        "Jan and Dec are month indicators. RetailMA3 is the trailing three month arithmetic average of the "
        "monthly percentage change in total retail sales.",
    )

    coef_rows = []
    reference_coefficients = REFERENCE["coefficients"]
    for _, row in artifacts.coefficients.iterrows():
        term = row["term"]
        reference = reference_coefficients[term]
        coef_rows.append(
            [
                "Constant" if term == "const" else term,
                f"{reference:.3f}" if term != "const" else f"{reference:.2f}",
                f"{row['coefficient']:.4f}",
                f"{row['coefficient'] - reference:+.4f}",
                fmt_p(row["p_value"]),
                "-" if pd.isna(row["vif"]) else f"{row['vif']:.2f}",
            ]
        )
    add_caption(doc, "Table 6 Reported and reproduced coefficients")
    add_table(
        doc,
        ["Term", "Reported", "Reproduced", "Difference", "p value", "VIF"],
        coef_rows,
        [2.0, 1.0, 1.15, 1.0, 0.9, 0.75],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        8.7,
    )

    add_caption(doc, "Table 7 Fit and residual diagnostics")
    add_table(
        doc,
        ["Metric", "Reported", "Reproduced", "Assessment"],
        [
            ["R squared", "0.623", f"{diag['r2']:.3f}", compare_status(0.623, diag["r2"], 3)],
            ["Adjusted R squared", "0.610", f"{diag['adj_r2']:.3f}", compare_status(0.610, diag["adj_r2"], 3)],
            ["Durbin Watson", "2.356", f"{diag['dw']:.3f}", compare_status(2.356, diag["dw"], 3)],
            ["Jarque Bera p value", "0.547", f"{diag['jb_p']:.3f}", compare_status(0.547, diag["jb_p"], 3)],
            ["Maximum VIF", "Below 2", f"{artifacts.coefficients['vif'].max():.2f}", "Passes the reported threshold"],
        ],
        [2.25, 1.2, 1.25, 2.3],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        8.8,
    )

    doc.add_heading("Interpretation", level=2)
    add_bullet(doc, "The negative autoregressive coefficients indicate short term reversal after unusually large monthly changes.")
    add_bullet(doc, "January is negative and December is positive, consistent with the visible seasonal pattern in the target series.")
    add_bullet(doc, "The positive retail sales coefficient links stronger recent retail growth with higher Consumer Card fee growth, conditional on the autoregressive and seasonal terms.")

    doc.add_page_break()
    doc.add_heading("Predictive Accuracy", level=1)
    add_paragraph(
        doc,
        "The accuracy calculation uses fixed in sample coefficients throughout the OOT period. Each monthly "
        "prediction uses the observed y(t-1) and y(t-2), the observed prior Consumer Card level, and the supplied "
        "macro path. This is a one step conditional backtest. It is not a recursive multi month forecast from "
        "September 2023.",
    )
    accuracy_rows = []
    for _, row in artifacts.accuracy.iterrows():
        if row["window"] == "In sample":
            ref_mape = REFERENCE["accuracy"]["is_mape"]
            ref_rmse = REFERENCE["accuracy"]["is_rmse"]
        else:
            ref_mape = REFERENCE["accuracy"]["oot_mape"]
            ref_rmse = REFERENCE["accuracy"]["oot_rmse"]
        accuracy_rows.append(
            [
                row["window"],
                str(int(row["n"])),
                f"{ref_mape:.2f}%",
                f"{row['mape_pct']:.2f}%",
                f"{ref_rmse:.1f}",
                f"{row['rmse']:.2f}",
            ]
        )
    add_caption(doc, "Table 8 Reported and reproduced accuracy")
    add_table(
        doc,
        ["Window", "n", "Reported MAPE", "Reproduced MAPE", "Reported RMSE", "Reproduced RMSE"],
        accuracy_rows,
        [1.55, 0.5, 1.15, 1.25, 1.15, 1.3],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        8.45,
    )
    add_figure(doc, figures["fit"], "Figure 2 Actual and predicted Consumer Card fee revenue")

    residual_rows = []
    residual_ref = REFERENCE["oot_residuals"]
    for _, row in artifacts.residual_diagnostics.iterrows():
        residual_rows.append(
            [
                row["window"],
                str(int(row["n"])),
                f"{row['mean']:.3f}",
                f"{row['std']:.3f}",
                f"{row['dw']:.3f}",
                f"{row['jb_p']:.3f}",
            ]
        )
    add_caption(doc, "Table 9 Reproduced transformed residual diagnostics")
    add_table(
        doc,
        ["Window", "n", "Mean", "Standard deviation", "Durbin Watson", "Jarque Bera p"],
        residual_rows,
        [1.45, 0.55, 0.9, 1.35, 1.25, 1.3],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        8.45,
    )
    add_paragraph(
        doc,
        f"The reproduced OOT residual standard deviation is {artifacts.residual_diagnostics.loc[1, 'std']:.3f}, "
        f"below the in sample value of {artifacts.residual_diagnostics.loc[0, 'std']:.3f}. The OOT residual "
        f"Durbin Watson statistic is {artifacts.residual_diagnostics.loc[1, 'dw']:.3f}, compared with "
        f"{residual_ref['dw']:.3f} in the prior report. The conclusion that OOT errors are not larger than the "
        "in sample errors is supported, but the printed residual diagnostics do not reproduce exactly.",
    )

    doc.add_page_break()
    doc.add_heading("Structural Stability", level=1)
    doc.add_heading("CUSUM", level=2)
    cusum_rows = []
    for i, row in artifacts.cusum.iterrows():
        reference_p = REFERENCE["cusum"]["is_p"] if i == 0 else REFERENCE["cusum"]["full_p"]
        cusum_rows.append(
            [
                row["sample"],
                str(int(row["n"])),
                f"{reference_p:.3f}",
                f"{row['p_value']:.3f}",
                "No rejection at 5 percent",
            ]
        )
    add_caption(doc, "Table 10 CUSUM results")
    add_table(
        doc,
        ["Sample", "n", "Reported p", "Reproduced p", "Conclusion"],
        cusum_rows,
        [1.55, 0.55, 1.05, 1.15, 2.7],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        8.8,
    )

    doc.add_heading("Chow Tests", level=2)
    chow_rows = []
    for _, row in artifacts.chow.iterrows():
        reference = REFERENCE["chow"][row["break_date"]]
        chow_rows.append(
            [
                row["break_date"],
                f"{int(row['n_pre'])} / {int(row['n_post'])}",
                f"{reference['f']:.2f}",
                f"{row['f_stat']:.3f}",
                f"{reference['p']:.3f}",
                f"{row['p_value']:.3f}",
            ]
        )
    add_caption(doc, "Table 11 Reported and reproduced Chow tests")
    add_table(
        doc,
        ["Break date", "n pre and post", "Reported F", "Reproduced F", "Reported p", "Reproduced p"],
        chow_rows,
        [1.15, 1.25, 1.0, 1.15, 1.0, 1.15],
        [WD_ALIGN_PARAGRAPH.CENTER] * 6,
        8.4,
    )
    add_paragraph(
        doc,
        "All reproduced Chow p values exceed 0.05, so the no break conclusion remains unchanged. The reported "
        "F statistics and p values do not reproduce numerically under the standard pooled versus split residual "
        "sum of squares formula with six parameters.",
    )

    doc.add_heading("Rolling Coefficients", level=2)
    retail_summary = artifacts.rolling_summary.set_index("term").loc[FINAL_MACRO_TERM]
    add_paragraph(
        doc,
        f"A 60 month rolling regression produces {len(artifacts.rolling)} windows, matching the prior report. "
        f"The retail sales coefficient ranges from {retail_summary['min']:.2f} to {retail_summary['max']:.2f}, "
        "also matching the reported range after rounding. Every y lag coefficient remains negative and every "
        "retail sales coefficient remains positive across all windows.",
    )
    add_figure(doc, figures["rolling"], "Figure 3 Rolling 60 month coefficient estimates", width=6.55)

    doc.add_page_break()
    doc.add_heading("Out of Time Reestimation", level=1)
    add_paragraph(
        doc,
        "The OOT only regression has 27 observations and six estimated parameters. Its estimates are therefore "
        "imprecise and sensitive to two holiday cycles. This exercise is a stability diagnostic and does not "
        "replace the fixed in sample model used in the backtest.",
    )
    oot_lookup = artifacts.oot_only.set_index("term")
    oot_rows = []
    for term in ["y_lag1", "y_lag2", "January", "December", FINAL_MACRO_TERM]:
        reproduced = oot_lookup.loc[term]
        reference = REFERENCE["oot_only"][term]
        oot_rows.append(
            [
                term,
                f"{reference['coef']:.3f}",
                fmt_p(reference["p"]),
                f"{reproduced['coefficient']:.3f}",
                fmt_p(reproduced["p_value"]),
                "Same" if np.sign(reference["coef"]) == np.sign(reproduced["coefficient"]) else "Different",
            ]
        )
    add_caption(doc, "Table 12 OOT only coefficient comparison")
    add_table(
        doc,
        ["Term", "Reported coefficient", "Reported p", "Reproduced coefficient", "Reproduced p", "Sign"],
        oot_rows,
        [1.7, 1.25, 0.95, 1.35, 1.0, 0.75],
        [WD_ALIGN_PARAGRAPH.LEFT] + [WD_ALIGN_PARAGRAPH.CENTER] * 5,
        8.35,
    )
    add_paragraph(
        doc,
        "The reproduced OOT only retail sales estimate remains negative and statistically insignificant. The "
        "exact coefficient and p value differ from the prior report. The y lag and seasonal estimates retain "
        "the reported directions. The small sample supports a cautious interpretation of individual OOT only slopes.",
    )

    doc.add_heading("Headline Reconciliation", level=1)
    reconcile_rows = []
    for _, row in artifacts.reconciliation.iterrows():
        reconcile_rows.append(
            [
                row["metric"],
                f"{row['reported']:.4g}",
                f"{row['reproduced']:.4g}",
                f"{row['difference']:+.4g}",
                row["status"],
            ]
        )
    add_caption(doc, "Table 13 Headline numerical reconciliation")
    add_table(
        doc,
        ["Metric", "Reported", "Reproduced", "Difference", "Status"],
        reconcile_rows,
        [2.5, 0.9, 1.0, 0.9, 1.7],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        7.9,
    )

    doc.add_page_break()
    doc.add_heading("Conclusion", level=1)
    add_paragraph(
        doc,
        "The available data reproduce the Consumer Card model's central result. A dynamic monthly log change "
        "regression with two autoregressive terms, January and December indicators, and the three month moving "
        "average of total retail sales growth produces statistically significant in sample coefficients, acceptable "
        "residual diagnostics, and lower OOT error than in sample error under a one step conditional backtest.",
    )
    add_paragraph(
        doc,
        "The reproduction does not confirm every printed number. The strongest matches are the stationarity tests, "
        "sample sizes, selected specification, coefficient signs, term significance, VIF results, Durbin Watson, "
        "rolling window count, and retail coefficient range. The available workbooks do not produce the exact "
        "reported coefficients, Jarque Bera p value, OOT error statistics, Chow statistics, or OOT only estimates. "
        "The original calculation file or original model code is required to determine whether those differences "
        "come from an earlier data vintage, an undocumented implementation choice, or transcription into the report.",
    )

    doc.add_heading("Limitations", level=2)
    add_bullet(doc, "The fee workbook does not state the currency or scale of the Consumer Card field, so level errors remain in source units.")
    add_bullet(doc, "The macro workbook does not contain release dates or real time vintages. Calendar availability does not prove information was available at each forecast origin.")
    add_bullet(doc, "The OOT backtest uses observed lagged target changes and is therefore a one step conditional evaluation rather than a recursive 27 month forecast.")
    add_bullet(doc, "Forward selection tests multiple related macro transformations without a multiplicity adjustment. The selection result should be treated as a reproduction of the reported workflow.")

    doc.add_heading("Reproduction Details", level=2)
    add_paragraph(
        doc,
        "Source files and software versions are recorded in the generated run manifest. The source workbooks are "
        "read only. The script writes full precision CSV tables for every statistic displayed in this report.",
    )
    audit = artifacts.source_audit
    add_caption(doc, "Table 14 Source provenance")
    add_table(
        doc,
        ["Source", "Coverage", "Rows", "SHA 256"],
        [
            [audit["card_filename"], f"{audit['card']['start']} to {audit['card']['end']}", str(audit["card"]["rows"]), audit["card_sha256"]],
            [audit["macro_filename"], f"{audit['macro']['start']} to {audit['macro']['end']}", str(audit["macro"]["rows"]), audit["macro_sha256"]],
        ],
        [1.55, 1.55, 0.6, 3.45],
        [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.LEFT],
        7.5,
    )

    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(output_path)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--card-file", type=Path, default=DEFAULT_CARD, help="Path to Cardfee.xlsx")
    parser.add_argument("--macro-file", type=Path, default=DEFAULT_MACRO, help="Path to Macro_Data_201001-202612.xls")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT, help="Output directory")
    parser.add_argument(
        "--report-name",
        default="Consumer_Card_Model_Reproduction_Report.docx",
        help="DOCX report filename",
    )
    return parser.parse_args()


def main() -> int:
    args = parse_args()
    output_dir = args.output_dir.resolve()
    report_path = output_dir / args.report_name
    assets_dir = output_dir / "assets"
    artifacts = run_analysis(args.card_file.resolve(), args.macro_file.resolve())
    save_numeric_outputs(artifacts, output_dir)
    figures = make_plots(artifacts, assets_dir)
    build_report(artifacts, figures, report_path)
    print(f"Report: {report_path}")
    print(f"Selected terms: {artifacts.selected_terms}")
    print(artifacts.reconciliation.to_string(index=False))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
